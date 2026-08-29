
from flask import Blueprint, render_template, request, jsonify, send_file
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.database import db
from app.models import (
    Bill,
    BillItem,
    Deposit,
    BillRefund,
    Patient
)


billing_reports_bp = Blueprint(
    "billing_reports",
    __name__,
    url_prefix="/billing-reports"
)


# =========================================================
# HELPER - GET REPORT FILTERS
# =========================================================

def get_report_filters():

    return {
        "start_date": request.args.get("start_date", "").strip(),
        "end_date": request.args.get("end_date", "").strip(),
        "hospital_no": request.args.get("hospital_no", "").strip(),
        "bill_no": request.args.get("bill_no", "").strip(),
        "patient_name": request.args.get("patient_name", "").strip(),
        "payment_type": request.args.get("payment_type", "").strip()
    }


# =========================================================
# HELPER - DATE RANGE
# =========================================================

def get_date_range(start_date, end_date):

    start = None
    end = None

    if start_date:
        start = datetime.strptime(
            start_date,
            "%Y-%m-%d"
        )

    if end_date:
        end = (
            datetime.strptime(
                end_date,
                "%Y-%m-%d"
            )
            + timedelta(days=1)
        )

    return start, end


# =========================================================
# BILLING REPORT PAGE
# =========================================================

@billing_reports_bp.route("/")
def billing_reports():

    return render_template(
        "reports/billing_reports.html"
    )


# =========================================================
# BILLING REPORT DATA
# =========================================================

@billing_reports_bp.route("/data")
def billing_report_data():

    try:

        filters = get_report_filters()

        start_date = filters["start_date"]
        end_date = filters["end_date"]
        hospital_no = filters["hospital_no"]
        bill_no = filters["bill_no"]
        patient_name = filters["patient_name"]
        payment_type = filters["payment_type"]

        start, end = get_date_range(
            start_date,
            end_date
        )

        # =================================================
        # BILLING
        # =================================================

        query = (
            db.session.query(
                Bill,
                Patient
            )
            .outerjoin(
                Patient,
                Patient.id == Bill.patient_id
            )
        )

        # -------------------------------------------------
        # BILL DATE
        # -------------------------------------------------

        if start:
            query = query.filter(
                Bill.bill_date >= start
            )

        if end:
            query = query.filter(
                Bill.bill_date < end
            )

        # -------------------------------------------------
        # HOSPITAL NUMBER
        # -------------------------------------------------

        if hospital_no:

            query = query.filter(
                Patient.patient_no.ilike(
                    f"%{hospital_no}%"
                )
            )

        # -------------------------------------------------
        # BILL NUMBER
        # -------------------------------------------------

        if bill_no:

            query = query.filter(
                Bill.bill_no.ilike(
                    f"%{bill_no}%"
                )
            )

        # -------------------------------------------------
        # PATIENT NAME
        # -------------------------------------------------

        if patient_name:

            query = query.filter(
                Patient.full_name.ilike(
                    f"%{patient_name}%"
                )
            )

        # -------------------------------------------------
        # PAYMENT TYPE
        # -------------------------------------------------

        if payment_type:

            query = query.filter(
                Bill.pay_type == payment_type
            )

        bills_result = query.order_by(
            Bill.bill_date.desc()
        ).all()

        bills = []

        gross_total = 0
        discount_total = 0
        net_total = 0

        for bill, patient in bills_result:

            subtotal = float(
                bill.subtotal or 0
            )

            discount = float(
                bill.discount or 0
            )

            total = float(
                bill.total or 0
            )

            gross_total += subtotal
            discount_total += discount
            net_total += total

            bills.append({

                "bill_no": bill.bill_no,

                "date": (
                    bill.bill_date.strftime("%Y-%m-%d")
                    if bill.bill_date
                    else ""
                ),

                "time": (
                    bill.bill_date.strftime("%I:%M %p")
                    if bill.bill_date
                    else ""
                ),

                "hospital_no": (
                    patient.patient_no
                    if patient
                    else ""
                ),

                "patient_name": (
                    patient.full_name
                    if patient
                    else ""
                ),

                "department": (
                    patient.department
                    if patient
                    else ""
                ),

                "subtotal": subtotal,

                "discount": discount,

                "total": total,

                "payment_type": (
                    bill.pay_type or ""
                ),

                "tender_amount": float(
                    bill.tender_amt or 0
                ),

                "return_amount": float(
                    bill.return_amt or 0
                ),

                "remarks": (
                    bill.remarks or ""
                )
            })

        # =================================================
        # REFUNDS
        # =================================================

        refund_query = (
            db.session.query(
                BillRefund
            )
        )

        # -------------------------------------------------
        # REFUND DATE
        # -------------------------------------------------

        if start:

            refund_query = refund_query.filter(
                BillRefund.created_at >= start
            )

        if end:

            refund_query = refund_query.filter(
                BillRefund.created_at < end
            )

        # -------------------------------------------------
        # REFUND PATIENT FILTER
        # -------------------------------------------------

        if hospital_no:

            refund_query = refund_query.filter(
                BillRefund.patient_no.ilike(
                    f"%{hospital_no}%"
                )
            )

        if patient_name:

            refund_query = refund_query.filter(
                BillRefund.patient_name.ilike(
                    f"%{patient_name}%"
                )
            )

        # -------------------------------------------------
        # REFUND BILL NUMBER
        # -------------------------------------------------

        if bill_no:

            refund_query = refund_query.filter(
                BillRefund.bill_no.ilike(
                    f"%{bill_no}%"
                )
            )

        refunds_result = refund_query.order_by(
            BillRefund.created_at.desc()
        ).all()

        refunds = []

        refund_total = 0

        for refund in refunds_result:

            amount = float(
                refund.refund_amount or 0
            )

            refund_total += amount

            refunds.append({

                "refund_no": refund.refund_no,

                "bill_no": refund.bill_no,

                "date": (
                    refund.created_at.strftime(
                        "%Y-%m-%d"
                    )
                    if refund.created_at
                    else ""
                ),

                "time": (
                    refund.created_at.strftime(
                        "%I:%M %p"
                    )
                    if refund.created_at
                    else ""
                ),

                "hospital_no": (
                    refund.patient_no or ""
                ),

                "patient_name": (
                    refund.patient_name or ""
                ),

                "bill_amount": float(
                    refund.bill_amount or 0
                ),

                "refund_amount": amount,

                "reason": (
                    refund.reason or ""
                )
            })

        # =================================================
        # DEPOSITS
        # =================================================

        deposit_query = (
            db.session.query(
                Deposit
            )
        )

        # -------------------------------------------------
        # DEPOSIT DATE
        # -------------------------------------------------

        if start:

            deposit_query = deposit_query.filter(
                Deposit.deposit_date >= start
            )

        if end:

            deposit_query = deposit_query.filter(
                Deposit.deposit_date < end
            )

        # -------------------------------------------------
        # DEPOSIT PATIENT FILTER
        # -------------------------------------------------

        if hospital_no:

            deposit_query = deposit_query.filter(
                Deposit.patient_no.ilike(
                    f"%{hospital_no}%"
                )
            )

        if patient_name:

            deposit_query = deposit_query.filter(
                Deposit.patient_name.ilike(
                    f"%{patient_name}%"
                )
            )

        deposits_result = deposit_query.order_by(
            Deposit.deposit_date.desc()
        ).all()

        deposits = []

        deposit_total = 0

        for deposit in deposits_result:

            amount = float(
                deposit.amount or 0
            )

            deposit_total += amount

            deposits.append({

                "deposit_no": (
                    deposit.receipt_no
                    or ""
                ),

                "date": (
                    deposit.deposit_date.strftime(
                        "%Y-%m-%d"
                    )
                    if deposit.deposit_date
                    else ""
                ),

                "time": (
                    deposit.deposit_date.strftime(
                        "%I:%M %p"
                    )
                    if deposit.deposit_date
                    else ""
                ),

                "hospital_no": (
                    deposit.patient_no or ""
                ),

                "patient_name": (
                    deposit.patient_name or ""
                ),

                "amount": amount,

                "remarks": (
                    deposit.remarks or ""
                )
            })

        # =================================================
        # NET REVENUE
        # =================================================

        net_revenue = (
            net_total
            - refund_total
        )

        # =================================================
        # RESPONSE
        # =================================================

        return jsonify({

            "success": True,

            "summary": {

                "total_bills": len(bills),

                "gross_billing": round(
                    gross_total,
                    2
                ),

                "discount": round(
                    discount_total,
                    2
                ),

                "net_billing": round(
                    net_total,
                    2
                ),

                "total_refund": round(
                    refund_total,
                    2
                ),

                "total_deposit": round(
                    deposit_total,
                    2
                ),

                "net_revenue": round(
                    net_revenue,
                    2
                )
            },

            "bills": bills,

            "refunds": refunds,

            "deposits": deposits

        })

    except Exception as e:

        db.session.rollback()

        return jsonify({

            "success": False,

            "message": str(e)

        }), 500

# =========================================================
# EXPORT BILLING REPORT TO EXCEL
# =========================================================

@main.route("/billing-reports/export-excel")
def export_billing_report_excel():

    try:

        # =================================================
        # FILTERS
        # =================================================

        start_date = request.args.get(
            "start_date",
            ""
        ).strip()

        end_date = request.args.get(
            "end_date",
            ""
        ).strip()

        hospital_no = request.args.get(
            "hospital_no",
            ""
        ).strip()

        bill_no = request.args.get(
            "bill_no",
            ""
        ).strip()

        patient_name = request.args.get(
            "patient_name",
            ""
        ).strip()

        payment_type = request.args.get(
            "payment_type",
            ""
        ).strip()

        # =================================================
        # DATE RANGE
        # =================================================

        start = None
        end = None

        if start_date:

            start = datetime.strptime(
                start_date,
                "%Y-%m-%d"
            )

        if end_date:

            end = (
                datetime.strptime(
                    end_date,
                    "%Y-%m-%d"
                )
                + timedelta(days=1)
            )

        # =================================================
        # BILLING QUERY
        # =================================================

        bill_query = (
            db.session.query(
                Bill,
                Patient
            )
            .outerjoin(
                Patient,
                Patient.id == Bill.patient_id
            )
        )

        if start:

            bill_query = bill_query.filter(
                Bill.bill_date >= start
            )

        if end:

            bill_query = bill_query.filter(
                Bill.bill_date < end
            )

        if hospital_no:

            bill_query = bill_query.filter(
                Patient.patient_no.ilike(
                    f"%{hospital_no}%"
                )
            )

        if bill_no:

            bill_query = bill_query.filter(
                Bill.bill_no.ilike(
                    f"%{bill_no}%"
                )
            )

        if patient_name:

            bill_query = bill_query.filter(
                Patient.full_name.ilike(
                    f"%{patient_name}%"
                )
            )

        if payment_type:

            bill_query = bill_query.filter(
                Bill.pay_type == payment_type
            )

        bills = bill_query.order_by(
            Bill.bill_date.desc()
        ).all()

        # =================================================
        # REFUND QUERY
        # =================================================

        refund_query = db.session.query(
            BillRefund
        )

        if start:

            refund_query = refund_query.filter(
                BillRefund.created_at >= start
            )

        if end:

            refund_query = refund_query.filter(
                BillRefund.created_at < end
            )

        if hospital_no:

            refund_query = refund_query.filter(
                BillRefund.patient_no.ilike(
                    f"%{hospital_no}%"
                )
            )

        if patient_name:

            refund_query = refund_query.filter(
                BillRefund.patient_name.ilike(
                    f"%{patient_name}%"
                )
            )

        if bill_no:

            refund_query = refund_query.filter(
                BillRefund.bill_no.ilike(
                    f"%{bill_no}%"
                )
            )

        refunds = refund_query.order_by(
            BillRefund.created_at.desc()
        ).all()

        # =================================================
        # DEPOSIT QUERY
        # =================================================

        deposit_query = db.session.query(
            Deposit
        )

        if start:

            deposit_query = deposit_query.filter(
                Deposit.deposit_date >= start
            )

        if end:

            deposit_query = deposit_query.filter(
                Deposit.deposit_date < end
            )

        if hospital_no:

            deposit_query = deposit_query.filter(
                Deposit.patient_no.ilike(
                    f"%{hospital_no}%"
                )
            )

        if patient_name:

            deposit_query = deposit_query.filter(
                Deposit.patient_name.ilike(
                    f"%{patient_name}%"
                )
            )

        deposits = deposit_query.order_by(
            Deposit.deposit_date.desc()
        ).all()

        # =================================================
        # CREATE EXCEL
        # =================================================

        workbook = Workbook()

        # =================================================
        # BILLING SHEET
        # =================================================

        billing_sheet = workbook.active
        billing_sheet.title = "Billing Transactions"

        billing_headers = [
            "SN",
            "Bill No.",
            "Date",
            "Time",
            "Hospital No.",
            "Patient Name",
            "Department",
            "Subtotal",
            "Discount",
            "Total",
            "Payment Type",
            "Tender Amount",
            "Return Amount",
            "Remarks"
        ]

        billing_sheet.append(
            billing_headers
        )

        for cell in billing_sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        for index, (bill, patient) in enumerate(
            bills,
            start=1
        ):

            billing_sheet.append([

                index,

                bill.bill_no,

                (
                    bill.bill_date.strftime(
                        "%Y-%m-%d"
                    )
                    if bill.bill_date
                    else ""
                ),

                (
                    bill.bill_date.strftime(
                        "%I:%M %p"
                    )
                    if bill.bill_date
                    else ""
                ),

                (
                    patient.patient_no
                    if patient
                    else ""
                ),

                (
                    patient.full_name
                    if patient
                    else ""
                ),

                (
                    patient.department
                    if patient
                    else ""
                ),

                float(
                    bill.subtotal or 0
                ),

                float(
                    bill.discount or 0
                ),

                float(
                    bill.total or 0
                ),

                bill.pay_type or "",

                float(
                    bill.tender_amt or 0
                ),

                float(
                    bill.return_amt or 0
                ),

                bill.remarks or ""

            ])

        # =================================================
        # REFUND SHEET
        # =================================================

        refund_sheet = workbook.create_sheet(
            "Refund Transactions"
        )

        refund_headers = [
            "SN",
            "Refund No.",
            "Bill No.",
            "Date",
            "Time",
            "Hospital No.",
            "Patient Name",
            "Bill Amount",
            "Refund Amount",
            "Reason"
        ]

        refund_sheet.append(
            refund_headers
        )

        for cell in refund_sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        for index, refund in enumerate(
            refunds,
            start=1
        ):

            refund_sheet.append([

                index,

                refund.refund_no,

                refund.bill_no,

                (
                    refund.created_at.strftime(
                        "%Y-%m-%d"
                    )
                    if refund.created_at
                    else ""
                ),

                (
                    refund.created_at.strftime(
                        "%I:%M %p"
                    )
                    if refund.created_at
                    else ""
                ),

                refund.patient_no or "",

                refund.patient_name or "",

                float(
                    refund.bill_amount or 0
                ),

                float(
                    refund.refund_amount or 0
                ),

                refund.reason or ""

            ])

        # =================================================
        # DEPOSIT SHEET
        # =================================================

        deposit_sheet = workbook.create_sheet(
            "Deposit Transactions"
        )

        deposit_headers = [
            "SN",
            "Receipt No.",
            "Date",
            "Time",
            "Hospital No.",
            "Patient Name",
            "Amount",
            "Remarks"
        ]

        deposit_sheet.append(
            deposit_headers
        )

        for cell in deposit_sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        for index, deposit in enumerate(
            deposits,
            start=1
        ):

            deposit_sheet.append([

                index,

                deposit.receipt_no or "",

                (
                    deposit.deposit_date.strftime(
                        "%Y-%m-%d"
                    )
                    if deposit.deposit_date
                    else ""
                ),

                (
                    deposit.deposit_date.strftime(
                        "%I:%M %p"
                    )
                    if deposit.deposit_date
                    else ""
                ),

                deposit.patient_no or "",

                deposit.patient_name or "",

                float(
                    deposit.amount or 0
                ),

                deposit.remarks or ""

            ])

        # =================================================
        # COLUMN WIDTH
        # =================================================

        for sheet in workbook.worksheets:

            for column_cells in sheet.columns:

                max_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:

                    if cell.value is not None:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                sheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    40
                )

            sheet.freeze_panes = "A2"

        # =================================================
        # OUTPUT
        # =================================================

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        filename = (
            "billing_report_"
            + datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )
            + ".xlsx"
        )

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )
        )

    except Exception as e:

        db.session.rollback()

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500